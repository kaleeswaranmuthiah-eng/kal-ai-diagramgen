from pptx import Presentation
from pptx.util import Inches

def generate_ppt(mermaid):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    slide.shapes.title.text = "Architecture Diagram"

    # placeholder box instead of image rendering for simplicity
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
    tf = txBox.text_frame
    tf.text = "Mermaid Diagram Generated:\n\n" + mermaid[:1000]

    out = "data/output/architecture.pptx"
    prs.save(out)
    return out